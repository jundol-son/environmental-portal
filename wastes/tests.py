from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

from django.core import mail
from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.utils import timezone

from .models import (
    AllocationRule,
    AllocationRuleDetail,
    Settlement,
    SettlementEmail,
    SettlementSource,
    Vendor,
    VendorPrice,
)
from .settlement_services import (
    SettlementError,
    apply_manual_adjustment,
    allocate_business_units,
    calculate_vendors,
    compare_sources,
    confirm_settlement,
    largest_remainder,
    mark_report_generated,
    reopen_settlement,
    settlement_period,
    validate_settlement,
)


def detail(name, rate, order):
    return SimpleNamespace(
        business_unit=name,
        allocation_rate=Decimal(str(rate)),
        sort_order=order,
    )


class AmountAllocationTests(TestCase):
    def test_settlement_period(self):
        self.assertEqual(
            settlement_period(date(2026, 9, 1)),
            (date(2026, 8, 21), date(2026, 9, 20)),
        )

    def test_largest_remainder_preserves_every_total(self):
        cases = [
            (1, [Decimal('33.33'), Decimal('33.33'), Decimal('33.34')]),
            (10, [Decimal('70'), Decimal('15'), Decimal('15')]),
            (101, [Decimal('90'), Decimal('10')]),
            (10001, [Decimal('70'), Decimal('15'), Decimal('15')]),
            (999999999999, [Decimal('70'), Decimal('15'), Decimal('15')]),
        ]
        for total, rates in cases:
            details = [
                detail(f'BU-{index}', rate, index)
                for index, rate in enumerate(rates)
            ]
            rows = largest_remainder(total, details)
            self.assertEqual(sum(row['final'] for row in rows), total)

    def test_tie_breaker_uses_rate_then_sort_order(self):
        rows = largest_remainder(1, [
            detail('later', 50, 2),
            detail('first', 50, 1),
        ])
        result = {row['detail'].business_unit: row['final'] for row in rows}
        self.assertEqual(result, {'later': 0, 'first': 1})


class SettlementWorkflowTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            username='settlement-admin',
            email='admin@example.com',
            password='test-password',
        )
        self.rule = AllocationRule.objects.create(
            rule_code='RATE_01',
            rule_name='70/15/15',
            valid_from=date(2026, 1, 1),
        )
        AllocationRuleDetail.objects.bulk_create([
            AllocationRuleDetail(
                rule=self.rule, business_unit='Foundry',
                allocation_rate=Decimal('70'), sort_order=1,
            ),
            AllocationRuleDetail(
                rule=self.rule, business_unit='연구소',
                allocation_rate=Decimal('15'), sort_order=2,
            ),
            AllocationRuleDetail(
                rule=self.rule, business_unit='CSS',
                allocation_rate=Decimal('15'), sort_order=3,
            ),
        ])
        self.transport = Vendor.objects.create(
            vendor_name='운반사',
            vendor_type=Vendor.Type.TRANSPORT,
            allocation_rule=self.rule,
        )
        self.disposal = Vendor.objects.create(
            vendor_name='처리사',
            vendor_type=Vendor.Type.DISPOSAL,
            allocation_rule=self.rule,
        )
        for vendor, price_type, price in (
            (self.transport, VendorPrice.Type.TRANSPORT, '10.0000'),
            (self.disposal, VendorPrice.Type.DISPOSAL, '20.0000'),
        ):
            VendorPrice.objects.create(
                vendor=vendor,
                waste_type='폐유',
                price_type=price_type,
                unit_price=Decimal(price),
                unit='kg',
                valid_from=date(2026, 1, 1),
            )
        self.settlement = Settlement.objects.create(
            settlement_month=date(2026, 9, 1),
            period_start=date(2026, 8, 21),
            period_end=date(2026, 9, 20),
            status=Settlement.Status.SOURCE_LOADED,
            created_by=self.user,
        )
        SettlementSource.objects.create(
            settlement=self.settlement,
            source_id='SOURCE-1',
            weigh_date=timezone.now(),
            waste_type='폐유',
            transport_company='운반사',
            disposal_company='처리사',
            weight=Decimal('10.001'),
            weight_unit='kg',
        )

    def test_complete_workflow_keeps_supply_vat_and_total_balanced(self):
        calculate_vendors(self.settlement, self.user)
        allocate_business_units(self.settlement, self.user)
        self.assertTrue(validate_settlement(self.settlement, self.user))
        self.settlement.refresh_from_db()
        self.assertEqual(self.settlement.status, Settlement.Status.VALIDATED)

        for result in self.settlement.vendor_results.all():
            allocations = result.allocations.all()
            self.assertEqual(
                sum(item.supply_amount for item in allocations),
                result.supply_amount,
            )
            self.assertEqual(
                sum(item.vat_amount for item in allocations),
                result.vat_amount,
            )
            self.assertEqual(
                sum(item.total_amount for item in allocations),
                result.total_amount,
            )

        confirm_settlement(self.settlement, self.user)
        self.settlement.refresh_from_db()
        self.assertEqual(self.settlement.status, Settlement.Status.CONFIRMED)

    def test_confirm_revalidates_the_saved_snapshot(self):
        calculate_vendors(self.settlement, self.user)
        allocate_business_units(self.settlement, self.user)
        self.assertTrue(validate_settlement(self.settlement, self.user))
        self.settlement.vendor_results.first().allocations.first().delete()

        with self.assertRaises(SettlementError):
            confirm_settlement(self.settlement, self.user)

        self.settlement.refresh_from_db()
        self.assertEqual(self.settlement.status, Settlement.Status.ALLOCATED)
        self.assertFalse(
            self.settlement.validations.get(code='ALLOCATION_SNAPSHOT').passed
        )

    def test_settlement_dashboard_requires_login(self):
        response = self.client.get('/wastes/settlements/')
        self.assertEqual(response.status_code, 302)
        self.client.force_login(self.user)
        response = self.client.get('/wastes/settlements/')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'DB에서 신규 데이터 불러오기')
        self.assertNotContains(response, '정산 생성 권한 필요')

    def test_settlement_history_filters_by_month_and_status(self):
        self.client.force_login(self.user)
        response = self.client.get(
            '/wastes/history/?month=2026-09&status=SOURCE_LOADED'
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '2026-09')
        self.assertContains(response, '1건')
        self.assertEqual(list(response.context['settlements']), [self.settlement])

    def test_creator_can_calculate_but_cannot_confirm_without_permission(self):
        creator = User.objects.create_user('settlement-creator', password='test')
        self.settlement.created_by = creator
        self.settlement.save(update_fields=['created_by'])
        self.client.force_login(creator)

        response = self.client.post(
            f'/wastes/settlements/{self.settlement.pk}/action/calculate/'
        )
        self.assertEqual(response.status_code, 302)
        self.settlement.refresh_from_db()
        self.assertEqual(
            self.settlement.status, Settlement.Status.VENDOR_CALCULATED
        )

        self.settlement.status = Settlement.Status.VALIDATED
        self.settlement.save(update_fields=['status'])
        self.client.post(
            f'/wastes/settlements/{self.settlement.pk}/action/confirm/'
        )
        self.settlement.refresh_from_db()
        self.assertEqual(self.settlement.status, Settlement.Status.VALIDATED)

    def test_authenticated_user_can_create_settlement_from_database(self):
        creator = User.objects.create_user('database-loader', password='test')
        self.client.force_login(creator)

        response = self.client.post(
            '/wastes/settlements/', {'settlement_month': '2026-10'}
        )

        settlement = Settlement.objects.get(settlement_month=date(2026, 10, 1))
        self.assertRedirects(
            response, f'/wastes/settlements/{settlement.pk}/'
        )
        self.assertEqual(settlement.created_by, creator)
        self.assertTrue(settlement.sources.exists())

    def test_practice_mode_uses_next_empty_month_and_fresh_mock_data(self):
        self.client.force_login(self.user)

        response = self.client.post('/wastes/settlements/', {
            'settlement_month': '2026-09',
            'mode': 'practice',
        })

        settlement = Settlement.objects.get(settlement_month=date(2026, 10, 1))
        self.assertRedirects(
            response, f'/wastes/settlements/{settlement.pk}/'
        )
        self.assertTrue(settlement.sources.exists())
        self.assertEqual(
            set(settlement.sources.values_list('source_reference', flat=True)),
            {'MockWeighingDataSource'},
        )

    def test_completed_settlement_can_be_reopened_for_rate_changes(self):
        calculate_vendors(self.settlement, self.user)
        allocate_business_units(self.settlement, self.user)
        validate_settlement(self.settlement, self.user)
        confirm_settlement(self.settlement, self.user)
        mark_report_generated(self.settlement, self.user)

        reopen_settlement(self.settlement, self.user, '분배율 변경')

        self.settlement.refresh_from_db()
        self.assertEqual(self.settlement.status, Settlement.Status.ALLOCATED)
        self.assertIsNone(self.settlement.confirmed_at)
        self.assertFalse(self.settlement.validations.exists())
        self.assertTrue(self.settlement.audit_logs.filter(
            action='REOPEN', reason='분배율 변경'
        ).exists())

    def test_manual_adjustment_preserves_vendor_total_and_requires_revalidation(self):
        calculate_vendors(self.settlement, self.user)
        allocate_business_units(self.settlement, self.user)
        validate_settlement(self.settlement, self.user)
        result = self.settlement.vendor_results.first()
        source, target = list(result.allocations.order_by('sort_order'))[:2]
        original_source = source.supply_amount
        original_target = target.supply_amount
        original_total = sum(
            row.supply_amount for row in result.allocations.all()
        )

        apply_manual_adjustment(
            self.settlement, source, target, 5, '사업부 추가 보정', self.user
        )

        source.refresh_from_db()
        target.refresh_from_db()
        self.settlement.refresh_from_db()
        self.assertEqual(source.supply_amount, original_source - 5)
        self.assertEqual(target.supply_amount, original_target + 5)
        self.assertEqual(
            sum(row.supply_amount for row in result.allocations.all()),
            original_total,
        )
        self.assertEqual(self.settlement.status, Settlement.Status.ALLOCATED)
        self.assertFalse(self.settlement.validations.exists())
        self.assertTrue(self.settlement.audit_logs.filter(
            action='MANUAL_ADJUST', reason='사업부 추가 보정'
        ).exists())

    @override_settings(
        WASTE_SETTLEMENT_EMAIL_BACKEND=(
            'django.core.mail.backends.locmem.EmailBackend'
        )
    )
    def test_email_preview_requires_approval_and_attaches_only_vendor_rows(self):
        from openpyxl import load_workbook

        for vendor in (self.transport, self.disposal):
            vendor.manager_name = f'{vendor.vendor_name} 담당자'
            vendor.manager_email = f'{vendor.pk}@example.com'
            vendor.mail_enabled = True
            vendor.save(update_fields=[
                'manager_name', 'manager_email', 'mail_enabled', 'updated_at'
            ])
        calculate_vendors(self.settlement, self.user)
        allocate_business_units(self.settlement, self.user)
        validate_settlement(self.settlement, self.user)
        confirm_settlement(self.settlement, self.user)
        mark_report_generated(self.settlement, self.user)
        self.client.force_login(self.user)

        self.client.post(
            f'/wastes/settlements/{self.settlement.pk}/action/prepare_emails/'
        )
        drafts = list(self.settlement.emails.order_by('vendor_name'))
        self.assertEqual(len(drafts), 2)

        send_url = (
            f'/wastes/settlements/{self.settlement.pk}/action/send_email/'
        )
        self.client.post(send_url, {'email_id': drafts[0].pk})
        drafts[0].refresh_from_db()
        self.assertEqual(drafts[0].status, SettlementEmail.Status.DRAFT)
        self.assertEqual(len(mail.outbox), 0)

        for draft in drafts:
            self.client.post(send_url, {
                'email_id': draft.pk,
                'approved': 'yes',
            })
        self.settlement.refresh_from_db()
        self.assertEqual(self.settlement.status, Settlement.Status.EMAIL_SENT)
        self.assertEqual(len(mail.outbox), 2)
        workbook = load_workbook(BytesIO(mail.outbox[0].attachments[0][1]))
        vendor_names = {
            row[0].value for row in workbook.active.iter_rows(min_row=2)
        }
        self.assertEqual(vendor_names, {drafts[0].vendor_name})
        self.assertEqual(
            self.settlement.audit_logs.filter(action='SEND_EMAIL').count(), 2
        )

    def test_vendor_can_select_a_different_allocation_rule(self):
        alternate = AllocationRule.objects.create(
            rule_code='RATE_02',
            rule_name='50/50',
            valid_from=date(2026, 1, 1),
        )
        AllocationRuleDetail.objects.bulk_create([
            AllocationRuleDetail(
                rule=alternate, business_unit='Foundry',
                allocation_rate=Decimal('50'), sort_order=1,
            ),
            AllocationRuleDetail(
                rule=alternate, business_unit='CSS',
                allocation_rate=Decimal('50'), sort_order=2,
            ),
        ])
        calculate_vendors(self.settlement, self.user)
        allocate_business_units(self.settlement, self.user)
        self.client.force_login(self.user)

        response = self.client.post(
            f'/wastes/settlements/{self.settlement.pk}/action/set_rule/',
            {'vendor_id': self.disposal.pk, 'allocation_rule': alternate.pk},
        )

        self.assertRedirects(
            response,
            f'/wastes/settlements/{self.settlement.pk}/?step=3',
        )
        self.disposal.refresh_from_db()
        self.settlement.refresh_from_db()
        self.assertEqual(self.disposal.allocation_rule, alternate)
        self.assertEqual(self.settlement.status, Settlement.Status.VENDOR_CALCULATED)
        self.assertFalse(self.settlement.vendor_results.first().allocations.exists())

        allocate_business_units(self.settlement, self.user)
        disposal_result = self.settlement.vendor_results.get(
            vendor=self.disposal
        )
        self.assertEqual(
            set(disposal_result.allocations.values_list('allocation_rate', flat=True)),
            {Decimal('50')},
        )

        detail = self.client.get(
            f'/wastes/settlements/{self.settlement.pk}/?step=3'
        )
        self.assertContains(detail, '업체별 분배율 설정 및 사업부 배분')
        self.assertContains(detail, 'RATE_02 - 50/50')

    def test_allocation_rule_management_can_edit_details(self):
        self.client.force_login(self.user)
        response = self.client.get(f'/wastes/allocations/?edit={self.rule.pk}')
        self.assertContains(response, 'Foundry|70.0000|1')
        self.assertContains(response, 'name="rule_id"')

        response = self.client.post('/wastes/allocations/', {
            'rule_id': self.rule.pk,
            'rule_code': self.rule.rule_code,
            'rule_name': '수정된 60/20/20',
            'valid_from': '2026-01-01',
            'valid_to': '',
            'active': 'on',
            'details_text': 'Foundry|60|1\n연구소|20|2\nCSS|20|3',
        })
        self.assertRedirects(response, '/wastes/allocations/')
        self.rule.refresh_from_db()
        self.assertEqual(self.rule.rule_name, '수정된 60/20/20')
        self.assertEqual(
            list(self.rule.details.values_list('allocation_rate', flat=True)),
            [Decimal('60'), Decimal('20'), Decimal('20')],
        )

    def test_allocation_snapshot_survives_master_rule_edit(self):
        calculate_vendors(self.settlement, self.user)
        allocate_business_units(self.settlement, self.user)
        result = self.settlement.vendor_results.first()
        self.assertEqual(result.allocation_rule_code, 'RATE_01')
        first_detail = self.rule.details.get(business_unit='Foundry')
        second_detail = self.rule.details.get(business_unit='연구소')
        first_detail.allocation_rate = Decimal('60')
        second_detail.allocation_rate = Decimal('25')
        first_detail.save(update_fields=['allocation_rate'])
        second_detail.save(update_fields=['allocation_rate'])

        self.assertTrue(validate_settlement(self.settlement, self.user))
        self.assertTrue(
            self.settlement.validations.get(code='ALLOCATION_SNAPSHOT').passed
        )
        self.assertEqual(
            set(result.allocations.values_list('allocation_rate', flat=True)),
            {Decimal('70'), Decimal('15')},
        )

    def test_source_comparison_does_not_overwrite_snapshot(self):
        class ChangedSource:
            def get_weighing_data(self, start_date, end_date):
                return [{
                    'source_id': 'SOURCE-2',
                    'weigh_date': timezone.now(),
                    'waste_type': '폐산',
                    'transport_company': '운반사',
                    'disposal_company': '처리사',
                    'weight': Decimal('20'),
                    'weight_unit': 'kg',
                }], False

        diff = compare_sources(self.settlement, ChangedSource())
        self.assertEqual(diff['added'], 1)
        self.assertEqual(diff['deleted'], 1)
        self.assertEqual(self.settlement.sources.count(), 1)

    def test_settlement_pages_report_and_excel_render(self):
        calculate_vendors(self.settlement, self.user)
        allocate_business_units(self.settlement, self.user)
        validate_settlement(self.settlement, self.user)
        confirm_settlement(self.settlement, self.user)
        mark_report_generated(self.settlement, self.user)

        self.client.force_login(self.user)
        paths = [
            f'/wastes/settlements/{self.settlement.pk}/',
            f'/wastes/settlements/{self.settlement.pk}/report/',
            '/wastes/history/',
            '/wastes/vendors/',
            '/wastes/prices/',
            '/wastes/allocations/',
            '/wastes/reports/',
            '/wastes/system/',
        ]
        for path in paths:
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 200)

        detail_response = self.client.get(
            f'/wastes/settlements/{self.settlement.pk}/'
        )
        self.assertContains(detail_response, 'data-bs-toggle="pill"', count=6)
        self.assertContains(detail_response, 'id="step-6"')
        self.assertContains(detail_response, '완료 정산 다시 열기')

        response = self.client.get(
            f'/wastes/settlements/{self.settlement.pk}/excel/'
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
